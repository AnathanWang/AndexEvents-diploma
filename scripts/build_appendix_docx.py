#!/usr/bin/env python3
"""Fill appendix .docx files (Courier 9pt, full source files)."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"

sys.path.insert(0, str(ROOT / "scripts"))
from generate_appendix_code_docs import CLIENT_FILES, SERVER_FILES  # noqa: E402


def add_caption(doc: Document, rel_path: str, caption: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"{rel_path} — {caption}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(6)


def add_code(doc: Document, code: str) -> None:
    for line in code.splitlines():
        p = doc.add_paragraph()
        run = p.add_run(line if line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        fmt = p.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0


def build_docx(out_path: Path, files: list[tuple[str, str]]) -> int:
    doc = Document()
    # узкие поля — больше кода на страницу
    for section in doc.sections:
        section.left_margin = Pt(42)
        section.right_margin = Pt(28)
        section.top_margin = Pt(42)
        section.bottom_margin = Pt(42)

    lines_total = 0
    for index, (rel_path, caption) in enumerate(files):
        path = ROOT / rel_path
        code = (
            path.read_text(encoding="utf-8", errors="replace").rstrip("\n")
            if path.exists()
            else f"// файл не найден: {rel_path}"
        )
        add_caption(doc, rel_path, caption)
        add_code(doc, code)
        lines_total += len(code.splitlines())
        if index < len(files) - 1:
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)

    doc.save(out_path)
    return lines_total


def main() -> None:
    client_path = DOCS / "appendix_client_code.docx"
    server_path = DOCS / "appendix_server_code.docx"
    c_lines = build_docx(client_path, CLIENT_FILES)
    s_lines = build_docx(server_path, SERVER_FILES)
    print(f"client: {client_path} ({len(CLIENT_FILES)} files, {c_lines} LOC, ~{c_lines / 42:.0f} pp)")
    print(f"server: {server_path} ({len(SERVER_FILES)} files, {s_lines} LOC, ~{s_lines / 42:.0f} pp)")


if __name__ == "__main__":
    main()
